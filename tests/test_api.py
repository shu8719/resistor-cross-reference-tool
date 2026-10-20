import io

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font

from app import app


def _xlsx(rows, post=None):
    """In-memory workbook from a list of rows (None leaves the cell empty)."""
    wb = Workbook()
    ws = wb.active
    for r in rows:
        ws.append(r)
    if post:
        post(ws)
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def _column(header, values):
    return _xlsx([[header]] + [[v] for v in values])


def _sheets(data):
    """{sheet name: {column header: [cell values]}} of a downloaded workbook."""
    wb = load_workbook(io.BytesIO(data))
    out = {}
    for ws in wb.worksheets:
        header, *rows = ws.iter_rows(values_only=True)
        out[ws.title] = {h: [r[i] for r in rows] for i, h in enumerate(header)}
    return out


@pytest.fixture
def client():
    app.config["TESTING"] = True
    return app.test_client()


@pytest.mark.parametrize(
    "pn, grade",
    [
        ("RN73R1JTTD1002B25", "◎"),
        ("RT0402BRD071KL", "◎"),
        ("TNPW04021K00BEEA", "◎"),
        ("ERA3ARB1002V", "◎"),
        ("RG1608N-1002-B-T5", "REVERSE"),
        ("CRCW04021K00FKEA", "×"),
        ("RT0402XXX999", "△"),
        ("rn73r-1j-ttd-1002-b-25", "◎"),  # normalisation
    ],
)
def test_sample_inputs(client, pn, grade):
    r = client.post("/analyze", json={"pn": pn})
    assert r.status_code == 200
    assert r.get_json()["grade"] == grade


@pytest.mark.parametrize(
    "pn",
    ["RN73R1JTTDRRRRB25", "RN73R1JTTD1R0RB25", "TNPW06031K0KBEEA", "RG1608N-RRR-B-T5"],
)
def test_malformed_resistance_does_not_crash(client, pn):
    r = client.post("/analyze", json={"pn": pn})
    assert r.status_code == 200
    assert r.get_json()["grade"] == "△"


def test_missing_pn_returns_400(client):
    r = client.post("/analyze", json={})
    assert r.status_code == 400
    assert "error" in r.get_json()


def test_bulk_roundtrip(client):
    buf = _column("Part Number", ["RN73R1JTTD1002B25", None, "CRCW04021K00FKEA"])
    r = client.post("/upload_bulk", data={"file": (buf, "in.xlsx")}, content_type="multipart/form-data")
    assert r.status_code == 200
    out = _sheets(r.data)["Results"]
    assert out["Match Grade"] == ["◎", "△", "×"]
    assert out["RG/RGV Proposal"][0] == "RG1608P-103-B-T5"


def test_bulk_without_file_returns_400(client):
    r = client.post("/upload_bulk", data={}, content_type="multipart/form-data")
    assert r.status_code == 400


def test_bulk_non_excel_returns_400(client):
    buf = io.BytesIO(b"Part Number\nRN73R1JTTD1002B25\n")
    r = client.post("/analyze_bulk", data={"file": (buf, "in.csv")}, content_type="multipart/form-data")
    assert r.status_code == 400
    assert r.get_json()["error"] == "Excel読込エラー"


def test_bulk_keeps_other_columns_and_ignores_trailing_blank_cells(client):
    def styled_tail(ws):
        ws.cell(row=30, column=1).font = Font(bold=True)  # formatted but empty

    rows = [["Qty", "MFR PN"], [100, "RN73R1JTTD1002B25"], [5, "RG1608P-103-B-T5"]]
    r = client.post(
        "/analyze_bulk", data={"file": (_xlsx(rows, styled_tail), "in.xlsx")}, content_type="multipart/form-data"
    ).get_json()
    assert r["pn_column"] == "MFR PN"
    assert r["row_count"] == 2

    r = client.post(
        "/upload_bulk", data={"file": (_xlsx(rows, styled_tail), "in.xlsx")}, content_type="multipart/form-data"
    )
    out = _sheets(r.data)["Results"]
    assert out["Qty"] == [100, 5]
    assert out["Match Grade"] == ["◎", "REVERSE"]


def test_analyze_payload_explains_match(client):
    r = client.post("/analyze", json={"pn": "RN73R1JTTD1002B25"}).get_json()
    assert r["maker_series"] == "KOA RN73R"
    assert [s["kind"] for s in r["segments"]] == ["series", "size", "terminal", "packaging", "res", "tol", "tcr"]
    assert [s["state"] for s in r["steps"]] == ["ok"] * 5
    assert r["steps"][-1]["detail_jp"] == "RG1608P-103-B-T5"
    assert [c["key"] for c in r["compare"]] == ["size", "res", "tol", "tcr"]
    assert r["compare"][0]["input"] == "1608（0603）"
    assert all(c["verdict"] in ("match", "equivalent") for c in r["compare"])


def test_analyze_payload_explains_out_of_range_with_retry(client):
    r = client.post("/analyze", json={"pn": "RN73R1JTTD1002F25"}).get_json()
    assert r["grade"] == "！"
    assert r["reason"]["title_jp"] == "公差が対応範囲外です"
    checks = {c["key"]: c for c in r["checks"]}
    assert checks["tol"]["ok"] is False and checks["size"]["ok"] is True
    assert [s["state"] for s in r["steps"]] == ["ok", "ok", "ok", "ng", "skip"]
    retry = {o["tol"]: o["pn"] for o in r["retry_options"]}
    assert retry == {0.5: "RN73R1JTTD1002D25", 0.1: "RN73R1JTTD1002B25"}


def test_analyze_payload_reverse_rows(client):
    r = client.post("/analyze", json={"pn": "RG1608P-103-B-T5"}).get_json()
    assert r["grade"] == "REVERSE"
    assert r["display_pn"] == "RG1608P-103-B-T5"
    assert [row["maker"] for row in r["competitor_rows"]] == ["KOA", "YAGEO", "Vishay", "Panasonic"]
    assert r["steps"][1]["state"] == "reverse"


def test_analyze_payload_out_of_scope_and_unknown(client):
    r = client.post("/analyze", json={"pn": "CRCW04021K00FKEA"}).get_json()
    assert r["grade"] == "×" and r["maker_series"] == "Vishay CRCW"
    assert r["steps"][1]["state"] == "ng"
    r = client.post("/analyze", json={"pn": "RT0402XXX999"}).get_json()
    assert r["grade"] == "△" and r["reason"]["title_jp"] == "品番の形式を判定できません"


def test_bulk_json(client):
    buf = _column("MFR PN", ["RN73R1JTTD1002B25", None, "RG1608P-103-B-T5", "RC0603FR-0710KL", "ABC-1234-XYZ"])
    r = client.post("/analyze_bulk", data={"file": (buf, "parts.xlsx")}, content_type="multipart/form-data")
    assert r.status_code == 200
    body = r.get_json()
    assert body["filename"] == "parts.xlsx"
    assert body["pn_column"] == "MFR PN"
    assert body["row_count"] == 5
    assert [row["row"] for row in body["rows"]] == [2, 3, 4, 5, 6]
    assert [row["grade"] for row in body["rows"]] == ["◎", "△", "REVERSE", "×", "△"]
    assert body["rows"][0]["proposal"] == "RG1608P-103-B-T5"
    assert body["rows"][2]["competitors"]["KOA"] == "RN73R1JTTD1002B25"
    assert body["rows"][3]["maker_series"] == "YAGEO RC"
    assert body["rows"][4]["note_jp"] == "品番の形式を判定できません"


def test_bulk_json_short_reason_for_out_of_range(client):
    buf = _column("PN", ["RN73R1JTTD1002F25"])
    r = client.post("/analyze_bulk", data={"file": (buf, "in.xlsx")}, content_type="multipart/form-data").get_json()
    assert r["rows"][0]["grade"] == "！"
    assert r["rows"][0]["note_jp"] == "公差対象外（±1%）"


def test_bulk_excel_has_reverse_sheet(client):
    buf = _column("Part Number", ["RG1608P-103-B-T5", "RN73R1JTTD1002B25"])
    r = client.post("/upload_bulk", data={"file": (buf, "in.xlsx")}, content_type="multipart/form-data")
    sheets = _sheets(r.data)
    assert list(sheets) == ["Results", "Reverse Lookup"]
    assert sheets["Results"]["Match Grade"] == ["REVERSE", "◎"]
    rev = sheets["Reverse Lookup"]
    assert rev["Row"] == [2]
    assert rev["KOA"][0] == "RN73R1JTTD1002B25"


def test_bulk_and_single_format_specs_the_same_way(client):
    """Single-PN report and bulk rows share one formatter (±1%, not ±1.0%)."""
    pn = "RN73R1JTTD1002F25"
    single = client.post("/analyze", json={"pn": pn}).get_json()
    buf = _column("PN", [pn])
    row = client.post("/analyze_bulk", data={"file": (buf, "in.xlsx")}, content_type="multipart/form-data").get_json()["rows"][0]
    assert single["tol"] == row["tol"] == "±1%"
    assert (single["size"], single["res"], single["tcr"]) == (row["size"], row["res"], row["tcr"])
