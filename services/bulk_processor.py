import io
from typing import Any, BinaryIO, Dict, List, Optional, Tuple

from openpyxl import Workbook, load_workbook

from services.analyzer import GRADE_OK, GRADE_OUT_OF_RANGE, GRADE_OUT_OF_SCOPE, GRADE_REVERSE, analyze_pn
from services.explain import short_reason
from utils.format import fmt_res, fmt_size, fmt_tcr, fmt_tol

COMPETITOR_MAKERS = ("KOA", "YAGEO", "Vishay", "Panasonic")
RESULT_COLUMNS = ("Match Grade", "RG/RGV Proposal", "Note", "Size", "Res", "Tol", "TCR")


def _cell_value(cell) -> Any:
    if cell.data_type == "e":  # error cells such as #N/A count as empty
        return None
    v = cell.value
    if isinstance(v, float) and v.is_integer():
        return int(v)
    return None if v == "" else v


def read_table(stream: BinaryIO) -> Tuple[List[Any], List[List[Any]]]:
    """First sheet as (header, data rows). Trailing empty cells and rows are dropped; rows share one width."""
    wb = load_workbook(stream, read_only=True, data_only=True)
    try:
        ws = wb.worksheets[0]
        ws.reset_dimensions()  # the stored dimension can be missing or wrong
        grid = [[_cell_value(c) for c in r] for r in ws.iter_rows()]
    finally:
        wb.close()
    for r in grid:
        while r and r[-1] is None:
            r.pop()
    while grid and not grid[-1]:
        grid.pop()
    if not grid:
        return [], []
    width = max(len(r) for r in grid)
    grid = [r + [None] * (width - len(r)) for r in grid]
    return grid[0], grid[1:]


def cell_text(v: Any) -> str:
    return "" if v is None else str(v)


def pick_pn_column(header: List[Any]) -> Optional[int]:
    """Index of the part-number column: first one named like PN / PART, else the first column."""
    if not header:
        return None
    for i, c in enumerate(header):
        cu = cell_text(c).upper()
        if ("PN" in cu) or ("PART" in cu):
            return i
    return 0


def _row_note(a) -> tuple:
    if a.grade == GRADE_REVERSE:
        return "逆引き(別シート参照)", "Reverse lookup (see separate sheet)"
    if a.grade == GRADE_OUT_OF_SCOPE:
        return "厚膜抵抗のため対象外", "Thick film (out of scope)"
    if a.grade == GRADE_OUT_OF_RANGE:
        return short_reason(a)
    if a.grade == GRADE_OK:
        return a.note_jp, a.note_en
    return "品番の形式を判定できません", "Part number format not recognized"


def analyze_row(pn: Any) -> Dict[str, Any]:
    """Classify one part number for bulk output (Excel columns and JSON rows)."""
    pn = cell_text(pn).strip()
    a = analyze_pn(pn)
    p = a.parsed
    note_jp, note_en = _row_note(a)
    return {
        "pn": pn,
        "maker": p.maker,
        "series": p.series,
        "maker_series": p.maker_series,
        "grade": a.grade,
        "proposal": "対象外" if a.grade == GRADE_OUT_OF_SCOPE else a.proposal_pn,
        "note_jp": note_jp,
        "note_en": note_en,
        "size": fmt_size(p.size),
        "res": fmt_res(p.res_ohm),
        "tol": fmt_tol(p.tol_pct),
        "tcr": fmt_tcr(p.tcr_ppm),
        "competitors": a.competitors if a.grade == GRADE_REVERSE else None,
    }


def analyze_bulk_rows(header: List[Any], data: List[List[Any]]) -> List[Dict[str, Any]]:
    """JSON-friendly rows for the bulk results table. `row` is the Excel row number (header = 1)."""
    col = pick_pn_column(header)
    return [{"row": i + 2, **analyze_row(r[col] if col is not None else "")} for i, r in enumerate(data)]


def _write_sheet(ws, header: List[Any], rows: List[List[Any]]) -> None:
    ws.append(header)
    for r in rows:
        ws.append(r)


def build_bulk_excel_output(header: List[Any], data: List[List[Any]]) -> io.BytesIO:
    rows = analyze_bulk_rows(header, data)
    wb = Workbook()
    ws = wb.active
    ws.title = "Results"
    _write_sheet(
        ws,
        [*header, *RESULT_COLUMNS],
        [
            [*src, r["grade"], r["proposal"], r["note_jp"], r["size"], r["res"], r["tol"], r["tcr"]]
            for src, r in zip(data, rows)
        ],
    )
    # Competitor equivalents for every RG / RGV row.
    reverse = [
        [r["row"], r["pn"], *((r["competitors"] or {}).get(m, "") for m in COMPETITOR_MAKERS)]
        for r in rows
        if r["grade"] == GRADE_REVERSE
    ]
    if reverse:
        _write_sheet(wb.create_sheet("Reverse Lookup"), ["Row", "Part Number", *COMPETITOR_MAKERS], reverse)
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output
